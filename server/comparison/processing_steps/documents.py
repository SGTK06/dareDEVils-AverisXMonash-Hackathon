from pathlib import Path


def document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".csv"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".xlsx":
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None, header=None)
        return "\n".join(" | ".join(str(v) for v in row if pd.notna(v)) for frame in sheets.values() for row in frame.values)
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    raise ValueError(f"Unsupported attachment format: {path.name}")
